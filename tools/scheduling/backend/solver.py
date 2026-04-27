#!/usr/bin/env python3
"""
日翊客服排班最佳化求解器 v2 (LP)
輸入：Excel 檔案路徑（argv[1]），輸出 Excel 路徑（argv[2]）
"""
import pulp, openpyxl, sys, json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def run(xlsx_in, xlsx_out):
    wb_in = openpyxl.load_workbook(xlsx_in, data_only=True)
    ws_in = wb_in.worksheets[0]
    rows_raw    = list(ws_in.iter_rows(min_row=1, values_only=True))
    date_row    = rows_raw[4]; weekday_row = rows_raw[5]; special_row = rows_raw[6]
    date_cols   = [i for i,v in enumerate(date_row) if isinstance(v, datetime)]

    NAMES = ['陳芝容','李雅筠','胡巧宜','游鈞捷','張語軒','廖聲友','廖武志','羅思凱','曾雪惠','劉康正','郭信智']
    N=len(NAMES); D=len(date_cols); ZHANG=NAMES.index('張語軒')

    ALL_S  = ['07~16','07:30~16:30','08~17','08~17','08~17','09~18','09~18','10~19','11~20','13~22','14~23']
    MEET_S = ['07~16','07:30~16:30','08~17','08~17','08~17','09~18','09~18','10~19','10~19','11~20','14~23']
    WKND_S = ['07~16','09~18','11~20','14~23']
    OFF_W  = ['休','特','國','例休','生日假','補休']
    isOff   = lambda v: any(o in str(v) for o in OFF_W) if v else False
    isEmpty = lambda v: not v or str(v).strip()==''

    def day_type(wd,sp):
        if sp and '課會' in sp: return 'meeting'
        if wd in ['六','日']: return 'weekend'
        if wd=='一': return 'mon'
        return 'weekday'

    def get_pool(dtype,zw=False):
        if dtype=='meeting': p=list(MEET_S)
        elif dtype=='mon': p=list(ALL_S); p.remove('08~17'); p.remove('10~19')
        elif dtype=='weekday': p=list(ALL_S); p.remove('08~17'); p.remove('09~18'); p.remove('10~19')
        else: p=list(WKND_S)
        if zw and '08~17' in p: p.remove('08~17')
        return p

    NEED={'mon':9,'weekday':8,'weekend':4,'meeting':11}
    dtypes  =[day_type(str(weekday_row[ci] or '').strip(),str(special_row[ci] or '').strip()) for ci in date_cols]
    weekdays=[str(weekday_row[ci] or '').strip() for ci in date_cols]
    specials=[str(special_row[ci] or '').strip() for ci in date_cols]
    dates_s =[date_row[ci].strftime('%m/%d') for ci in date_cols]
    orig    =[[str(rows_raw[pi+7][ci] or '').strip() for ci in date_cols] for pi in range(N)]
    week_starts=[d for d in range(D) if weekdays[d]=='一']
    week_ends  =[week_starts[i+1]-1 if i+1<len(week_starts) else D-1 for i in range(len(week_starts))]

    # LP 模型
    prob=pulp.LpProblem("sched",pulp.LpMinimize)
    work=[[pulp.LpVariable(f"w_{i}_{d}",cat='Binary') for d in range(D)] for i in range(N)]

    for i in range(N):
        for d in range(D):
            v=orig[i][d]
            if isOff(v): prob+=work[i][d]==0
            elif not isEmpty(v): prob+=work[i][d]==1

    for d in range(D):
        if dtypes[d]=='weekend': prob+=work[ZHANG][d]==0

    for d in range(D):
        zhang_on=1 if dtypes[d]!='weekend' else 0
        prob+=pulp.lpSum(work[i][d] for i in range(N) if i!=ZHANG)==NEED[dtypes[d]]-zhang_on

    for i in range(N):
        if i==ZHANG: continue
        for ws_d,we_d in zip(week_starts,week_ends):
            prob+=pulp.lpSum(work[i][d] for d in range(ws_d,we_d+1))<=(we_d-ws_d+1)-2

    for i in range(N):
        for d in range(D-6):
            prob+=pulp.lpSum(work[i][d+k] for k in range(7))<=6

    non_zhang=[i for i in range(N) if i!=ZHANG]
    totals=[pulp.lpSum(work[i][d] for d in range(D)) for i in non_zhang]
    max_w=pulp.LpVariable("max_w",lowBound=0); min_w=pulp.LpVariable("min_w",lowBound=0)
    for t in totals: prob+=max_w>=t; prob+=min_w<=t
    prob+=max_w-min_w<=8

    zigzag=[]
    for i in range(N):
        if i==ZHANG: continue
        for d in range(1,D-1):
            z=pulp.LpVariable(f"z_{i}_{d}",cat='Binary')
            prob+=z<=work[i][d-1]; prob+=z<=1-work[i][d]; prob+=z<=work[i][d+1]
            prob+=z>=work[i][d-1]+(1-work[i][d])+work[i][d+1]-2
            zigzag.append(z)

    prob+=pulp.lpSum(zigzag)*5+(max_w-min_w)
    pulp.PULP_CBC_CMD(msg=0,timeLimit=180,gapRel=0.03).solve(prob)

    status=pulp.LpStatus[prob.status]
    if status not in ('Optimal','Feasible'):
        raise RuntimeError(f"求解失敗：{status}")

    # 提取結果
    result=[]
    for i in range(N):
        shifts=[]
        for d in range(D):
            v=orig[i][d]
            if isOff(v): shifts.append(v)
            elif not isEmpty(v): shifts.append(v)
            elif i==ZHANG: shifts.append('休' if dtypes[d]=='weekend' else '08~17')
            elif (pulp.value(work[i][d]) or 0)>0.5: shifts.append(None)
            else: shifts.append('休')
        result.append(shifts)

    # 班位分配
    for d in range(D):
        zw=result[ZHANG][d] not in ('休',) and not isOff(result[ZHANG][d] or '') and result[ZHANG][d] is not None
        pool=get_pool(dtypes[d],zw)
        for i in range(N):
            if i==ZHANG: continue
            v=result[i][d]
            if v and v!='休' and not isOff(v) and v is not None:
                if v in pool: pool.remove(v)
        for i in range(N):
            if i==ZHANG: continue
            if result[i][d] is None:
                result[i][d]=pool.pop(0) if pool else '休'
        if result[ZHANG][d] is None: result[ZHANG][d]='08~17'

    # 輸出 Excel
    THIN=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
    wb_out=Workbook(); ws_out=wb_out.active; ws_out.title='排班表'
    ws_out.freeze_panes='B4'

    HDR='2D3748'
    ws_out.cell(1,1,'姓名').font=Font(bold=True,color='FFFFFF',name='Arial',size=9)
    ws_out.cell(1,1).fill=PatternFill('solid',fgColor=HDR)
    ws_out.cell(1,1).alignment=Alignment(horizontal='center')
    for d,dt in enumerate(dates_s):
        c=ws_out.cell(1,d+2,dt)
        c.font=Font(bold=True,color='FFFFFF',name='Arial',size=9)
        c.fill=PatternFill('solid',fgColor=HDR); c.alignment=Alignment(horizontal='center')

    ws_out.cell(2,1,'星期').font=Font(bold=True,name='Arial',size=9)
    ws_out.cell(2,1).fill=PatternFill('solid',fgColor='EDF2F7')
    for d,wd in enumerate(weekdays):
        bg='EBF5FB' if dtypes[d]=='weekend' else ('FFF3CD' if dtypes[d]=='meeting' else 'EDF2F7')
        c=ws_out.cell(2,d+2,wd)
        c.font=Font(name='Arial',size=9); c.fill=PatternFill('solid',fgColor=bg)
        c.alignment=Alignment(horizontal='center')

    ws_out.cell(3,1,'備注').font=Font(bold=True,name='Arial',size=9)
    ws_out.cell(3,1).fill=PatternFill('solid',fgColor='EDF2F7')
    for d,sp in enumerate(specials):
        bg='FFF3CD' if '課會' in (sp or '') else ('FEF3C7' if sp else 'EDF2F7')
        c=ws_out.cell(3,d+2,sp or '')
        c.font=Font(name='Arial',size=9); c.fill=PatternFill('solid',fgColor=bg)
        c.alignment=Alignment(horizontal='center')

    for pi,name in enumerate(NAMES):
        row=pi+4
        ws_out.cell(row,1,name).font=Font(bold=True,name='Arial',size=9)
        ws_out.cell(row,1).alignment=Alignment(horizontal='left',vertical='center')
        for d in range(D):
            val=result[pi][d] or ''
            c=ws_out.cell(row,d+2,val)
            c.alignment=Alignment(horizontal='center',vertical='center'); c.border=THIN
            is_new=isEmpty(orig[pi][d]) and not isOff(val) and val and val!='休'
            if isOff(val):
                c.font=Font(bold=True,color='C53030',name='Arial',size=9)
                c.fill=PatternFill('solid',fgColor='FFC7CE')
            elif dtypes[d]=='meeting':
                c.font=Font(name='Arial',size=9); c.fill=PatternFill('solid',fgColor='FFF3CD')
            elif dtypes[d]=='weekend':
                c.font=Font(name='Arial',size=9); c.fill=PatternFill('solid',fgColor='EBF5FB')
            elif is_new:
                c.font=Font(name='Arial',size=9,color='15803D'); c.fill=PatternFill('solid',fgColor='C7EFCE')
            else:
                c.font=Font(name='Arial',size=9); c.fill=PatternFill('solid',fgColor='FFFFFF')

    ws_out.column_dimensions['A'].width=10
    for d in range(D): ws_out.column_dimensions[get_column_letter(d+2)].width=11
    for r in range(1,N+5): ws_out.row_dimensions[r].height=18

    # 統計
    ws_stat=wb_out.create_sheet('統計')
    ws_stat.append(['人員','總上班','總休假','週1休','週2休','週3休','週4休','週5休','最長連上'])
    for pi,name in enumerate(NAMES):
        tw=sum(1 for d in range(D) if not isOff(result[pi][d]) and result[pi][d] and result[pi][d]!='休')
        to=D-tw
        woffs=[sum(1 for d in range(ws_d,we_d+1) if isOff(result[pi][d])) for ws_d,we_d in zip(week_starts,week_ends)]
        cons=0; mc=0
        for d in range(D):
            if not isOff(result[pi][d]) and result[pi][d] and result[pi][d]!='休': cons+=1; mc=max(mc,cons)
            else: cons=0
        ws_stat.append([name,tw,to]+woffs+[mc])

    # 法規警告
    ws_warn=wb_out.create_sheet('法規警告')
    ws_warn.append(['類型','人員','說明']); warns=[]
    for pi,name in enumerate(NAMES):
        cons=0
        for d in range(D):
            if not isOff(result[pi][d]) and result[pi][d] and result[pi][d]!='休': cons+=1
            else: cons=0
            if cons>6: warns.append(['連班超標',name,f"{dates_s[d]}連上第{cons}天"])
        for wi,(ws_d,we_d) in enumerate(zip(week_starts,week_ends)):
            off=sum(1 for d in range(ws_d,we_d+1) if isOff(result[pi][d]))
            if off<2 and pi!=ZHANG: warns.append(['週休不足',name,f"第{wi+1}週{off}天"])
        zz=sum(1 for d in range(1,D-1)
               if not isOff(result[pi][d-1]) and result[pi][d-1] and result[pi][d-1]!='休'
               and isOff(result[pi][d])
               and not isOff(result[pi][d+1]) and result[pi][d+1] and result[pi][d+1]!='休')
        if zz>=2: warns.append(['碎班',name,f"{zz}次上一休一"])
    for w in warns: ws_warn.append(w)
    if not warns: ws_warn.append(['✓','全員','無法規問題'])

    wb_out.save(xlsx_out)
    return {'status':status, 'warnings':warns, 'n_warnings':len(warns)}

if __name__=='__main__':
    r=run(sys.argv[1], sys.argv[2])
    print(json.dumps(r, ensure_ascii=False))
